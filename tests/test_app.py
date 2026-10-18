# -*- coding: utf-8 -*-
# pyright: reportUninitializedInstanceVariable=false, reportAssignmentType=false
# pyright: reportMissingTypeArgument=false, reportUnknownArgumentType=false
# pyright: reportUnknownMemberType=false, reportAny=false, reportExplicitAny=false
"""
单元测试 - 个人工具箱
运行方式: python -m pytest tests/ -v
或: python tests/test_app.py
"""
import io
import sys
import os
import unittest
import tempfile
import sqlite3
from typing import override

# 将项目根目录加入路径
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import app
from flask.testing import FlaskClient

import app as app_mod
import routes.ledger as ledger_mod
from routes.utils import make_db

# ledger 模块真实的 DB 连接工厂（扩展测试替换后需要恢复）
_REAL_LEDGER_GET_DB = ledger_mod._get_db


class TestAppBasics(unittest.TestCase):
    """应用基础测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_index_200(self) -> None:
        r = self.client.get('/')
        self.assertEqual(r.status_code, 200)

    def test_api_version(self) -> None:
        r = self.client.get('/api/version')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('version', data)
        self.assertIn('build', data)

    def test_api_health(self) -> None:
        r = self.client.get('/api/ping')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertEqual(data.get('status'), 'ok')

    def test_nav_page(self) -> None:
        r = self.client.get('/nav')
        self.assertEqual(r.status_code, 200)

    def test_static_icon(self) -> None:
        r = self.client.get('/nav/icon-192.svg')
        self.assertIn(r.status_code, [200, 404])

    def test_404_unknown(self) -> None:
        r = self.client.get('/nonexistent-path-xyz')
        self.assertIn(r.status_code, [404, 302])


class TestRenqingAPI(unittest.TestCase):
    """人情礼金 API 测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_list_events(self) -> None:
        r = self.client.get('/api/renqing/events')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('success', data)

    def test_list_events_with_params(self) -> None:
        r = self.client.get('/api/renqing/events?page=1&page_size=5')
        self.assertEqual(r.status_code, 200)

    def test_stats(self) -> None:
        r = self.client.get('/api/renqing/stats')
        self.assertEqual(r.status_code, 200)

    def test_list_names(self) -> None:
        r = self.client.get('/api/renqing/names')
        self.assertEqual(r.status_code, 200)


class TestGPAAPI(unittest.TestCase):
    """绩点 API 测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_page_accessible(self) -> None:
        r = self.client.get('/gpa')
        self.assertEqual(r.status_code, 200)

    def test_list_courses(self) -> None:
        r = self.client.get('/api/gpa/courses')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('success', data)

    def test_get_stats(self) -> None:
        r = self.client.get('/api/gpa/stats')
        self.assertEqual(r.status_code, 200)


class TestHSGradesAPI(unittest.TestCase):
    """成绩 API 测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_page_accessible(self) -> None:
        r = self.client.get('/hsgrades')
        self.assertEqual(r.status_code, 200)

    def test_list_exams(self) -> None:
        r = self.client.get('/api/hsgrades/exams')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('success', data)


class TestAccountingAPI(unittest.TestCase):
    """记账 API 测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_page_accessible(self) -> None:
        r = self.client.get('/accounting')
        self.assertEqual(r.status_code, 200)

    def test_list_records(self) -> None:
        r = self.client.get('/api/accounting/records')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('success', data)

    def test_get_stats(self) -> None:
        r = self.client.get('/api/accounting/stats')
        self.assertEqual(r.status_code, 200)

    def test_list_categories(self) -> None:
        r = self.client.get('/api/accounting/categories')
        self.assertEqual(r.status_code, 200)


class TestCountdownAPI(unittest.TestCase):
    """倒计时 API 测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_page_accessible(self) -> None:
        r = self.client.get('/countdown')
        self.assertEqual(r.status_code, 200)

    def test_list_events(self) -> None:
        r = self.client.get('/api/countdown/events')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('success', data)


class TestLedgerAPI(unittest.TestCase):
    """专业记账会计 API 测试（需要登录会话）"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True
        with cls.client.session_transaction() as sess:  # type: ignore[attr-defined]
            sess['auth'] = True

    def test_page_accessible(self) -> None:
        r = self.client.get('/ledger')
        self.assertEqual(r.status_code, 200)
        self.assertIn('记账会计', r.get_data(as_text=True))

    def test_manifest(self) -> None:
        r = self.client.get('/ledger/manifest.json')
        self.assertEqual(r.status_code, 200)

    def test_list_accounts(self) -> None:
        r = self.client.get('/api/ledger/accounts')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('success', data)
        # 空库/已种子库都应返回数据数组
        self.assertIsInstance(data.get('data'), list)

    def test_dashboard(self) -> None:
        r = self.client.get('/api/ledger/dashboard')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertEqual(data.get('success'), True)
        self.assertIn('monthly', data)

    def test_trial(self) -> None:
        r = self.client.get('/api/ledger/trial?month=2026-08')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertIn('rows', data)

    def test_contacts(self) -> None:
        r = self.client.get('/api/ledger/contacts')
        self.assertEqual(r.status_code, 200)
        data: dict = r.get_json()
        self.assertEqual(data.get('success'), True)


class TestLedgerExtAPI(unittest.TestCase):
    """专业记账扩展 API 测试（辅助核算/凭证审核/红冲/模板/结转/结账）

    使用独立临时数据库，每个测试重建干净库，不触碰真实数据。
    """

    client: FlaskClient = None
    _tmp_dir: str = ''
    _names: set[str] = set()
    _seq: int = 0

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True
        with cls.client.session_transaction() as sess:  # type: ignore[attr-defined]
            sess['auth'] = True
        cls._tmp_dir = tempfile.mkdtemp(prefix='ledger_ext_')
        cls._names = {'测试客户', '测试客户2', '测试客户3', '测试应收-甲', '测试模板'}
        cls._seq = 0

    @classmethod
    @override
    def tearDownClass(cls) -> None:
        import gc
        import shutil
        gc.collect()
        shutil.rmtree(cls._tmp_dir, ignore_errors=True)
        ledger_mod._get_db = _REAL_LEDGER_GET_DB  # 恢复真实库连接，避免污染其它测试

    def _new_db(self) -> str:
        """为当前测试创建唯一临时库文件，替换连接工厂并初始化表结构"""
        import uuid
        self._seq += 1
        name = f'ledger_ext_{self._seq}_{uuid.uuid4().hex[:8]}.db'
        path = os.path.join(self._tmp_dir, name)
        ledger_mod._get_db = make_db(path)  # 替换连接工厂，隔离真实库
        ledger_mod._init_db_impl()
        return path

    def setUp(self) -> None:
        self._new_db()
        app_mod._rate_limits.clear()  # 避免速率限制干扰测试

    def _call(self, method: str, url: str, **kw):
        """清限流后调用 API"""
        app_mod._rate_limits.clear()
        r = getattr(self.client, method)(url, **kw)
        try:
            return r.status_code, r.get_json()
        except Exception:
            return r.status_code, None

    def _acct(self, code: str) -> dict:
        _, data = self._call('get', '/api/ledger/accounts')
        accts: list = data['data'] or []
        for a in accts:
            if a['code'] == code:
                return a
        raise AssertionError(f'科目 {code} 不存在')

    def _mk_dim(self, name: str) -> int:
        """创建辅助维度并回查 id（创建接口不返回 id）"""
        sc, d = self._call('post', '/api/ledger/aux/dims', json={'name': name})
        assert sc == 200 and d.get('success'), (name, sc, d)
        _, dims = self._call('get', '/api/ledger/aux/dims')
        dim = [x for x in dims['data'] if x['name'] == name][0]
        return dim['id']

    def _mk_item(self, dim_id: int, name: str) -> int:
        """创建辅助项并回查 id"""
        sc, d = self._call('post', '/api/ledger/aux/items',
                           json={'dim_id': dim_id, 'name': name})
        assert sc == 200 and d.get('success'), (name, sc, d)
        _, dims = self._call('get', '/api/ledger/aux/dims')
        dim = [x for x in dims['data'] if x['id'] == dim_id][0]
        item = [i for i in dim['items'] if i['name'] == name][0]
        return item['id']

    def _del_dim(self, dim_id: int) -> None:
        """删除维度（先删项）"""
        _, dims = self._call('get', '/api/ledger/aux/dims')
        dim = [x for x in dims['data'] if x['id'] == dim_id]
        if not dim:
            return
        for it in dim[0].get('items') or []:
            self._call('delete', f"/api/ledger/aux/items/{it['id']}")
        self._call('delete', f'/api/ledger/aux/dims/{dim_id}')

    def _acct_id(self, code: str) -> int:
        return self._acct(code)['id']

    def test_pc_pages(self) -> None:
        """桌面版页面与移动页面均可访问"""
        r = self.client.get('/ledger/pc')
        self.assertEqual(r.status_code, 200)
        self.assertIn('记账会计 · 财务桌面', r.get_data(as_text=True))
        r2 = self.client.get('/ledger')
        self.assertEqual(r2.status_code, 200)

    def test_aux_dims_and_items_crud(self) -> None:
        """辅助核算维度/项 创建、查询、删除"""
        sc, d = self._call('get', '/api/ledger/aux/dims')
        self.assertEqual(sc, 200)
        self.assertEqual(d['success'], True)
        # 创建维度并回查
        dim_id = self._mk_dim('测试客户')
        # 建两项
        item_a = self._mk_item(dim_id, '公司A')
        self._mk_item(dim_id, '公司B')
        # 查询含子项
        _, dims = self._call('get', '/api/ledger/aux/dims')
        dim = [x for x in dims['data'] if x['name'] == '测试客户'][0]
        self.assertEqual(len(dim['items']), 2)
        # 删除项
        sc, _ = self._call('delete', f'/api/ledger/aux/items/{item_a}')
        self.assertEqual(sc, 200)
        _, dims = self._call('get', '/api/ledger/aux/dims')
        dim = [x for x in dims['data'] if x['name'] == '测试客户'][0]
        self.assertEqual(len(dim['items']), 1)
        # 删除维度（先删剩余辅助项）
        self._del_dim(dim['id'])
        _, dims = self._call('get', '/api/ledger/aux/dims')
        self.assertNotIn('测试客户', [x['name'] for x in dims['data']])

    def test_account_bind_and_aux_opening(self) -> None:
        """科目绑定维度 + 下级科目 + 辅助期初读写"""
        # 准备维度与科目
        dim_id = self._mk_dim('测试客户2')
        item_a = self._mk_item(dim_id, '公司A')
        item_b = self._mk_item(dim_id, '公司B')
        # 绑定维度到银行存款 1002
        cash = self._acct('1002')
        sc, d = self._call('put', f"/api/ledger/accounts/{cash['id']}", json={
            'code': '1002', 'name': '银行存款', 'category': 'asset', 'parent_code': '',
            'aux_dims': [str(dim_id)], 'opening': 0, 'remark': ''})
        self.assertEqual(sc, 200)
        # 新建下级科目 112299（父 1122），绑维度
        sc, d = self._call('post', '/api/ledger/accounts', json={
            'code': '112299', 'name': '测试应收-甲', 'category': 'asset',
            'parent_code': '1122', 'aux_dims': [str(dim_id)], 'opening': 0, 'remark': ''})
        self.assertEqual(sc, 200)
        # 写入辅助期初并回读
        sc, d = self._call('put', '/api/ledger/aux/openings', json={
            'account_code': '112299',
            'rows': [{'item_id': item_a, 'amount': 1000.0}, {'item_id': item_b, 'amount': 500.0}]})
        self.assertEqual(sc, 200)
        sc, d = self._call('get', '/api/ledger/aux/openings?account_code=112299')
        self.assertEqual(d['success'], True)
        self.assertEqual(len(d['data']), 2)
        # 清理：解除绑定、删除科目与维度
        sub = self._acct('112299')
        self._call('put', '/api/ledger/aux/openings', json={'account_code': '112299', 'rows': []})
        self._call('delete', f"/api/ledger/accounts/{sub['id']}")
        cash = self._acct('1002')
        self._call('put', f"/api/ledger/accounts/{cash['id']}", json={
            'code': '1002', 'name': '银行存款', 'category': 'asset', 'parent_code': '',
            'aux_dims': [], 'opening': 0, 'remark': ''})
        self._del_dim(dim_id)

    def test_voucher_aux_audit_void_reverse(self) -> None:
        """凭证带辅助核算 + 审核/作废/恢复 + 红冲"""
        dim_id = self._mk_dim('测试客户3')
        item_a = self._mk_item(dim_id, '公司A')
        # 绑定维度到 1122 应收账款
        ar = self._acct('1122')
        sc, _ = self._call('put', f"/api/ledger/accounts/{ar['id']}", json={
            'code': '1122', 'name': ar['name'], 'category': ar['category'],
            'parent_code': ar['parent_code'], 'aux_dims': [str(dim_id)],
            'opening': ar.get('opening') or 0, 'remark': ar.get('remark') or ''})
        self.assertEqual(sc, 200)
        # 正确维度凭证可过
        sc, d = self._call('post', '/api/ledger/vouchers', json={
            'vdate': '2099-01-05', 'summary': '带辅助凭证', 'attachments': 1,
            'entries': [
                {'summary': 'A公司挂账', 'account_code': '1122', 'debit': 1000, 'credit': 0,
                 'aux': [{'dim_id': dim_id, 'item_id': item_a}]},
                {'summary': '收现金', 'account_code': '1001', 'debit': 0, 'credit': 1000},
            ]})
        self.assertEqual(sc, 200)
        self.assertEqual(d['success'], True)
        vid: int = d['id']
        # 维度不存在 → 拒绝
        sc, bad = self._call('post', '/api/ledger/vouchers', json={
            'vdate': '2099-01-06', 'summary': '坏维度',
            'entries': [
                {'summary': '错维度', 'account_code': '1122', 'debit': 10, 'credit': 0,
                 'aux': [{'dim_id': 99999, 'item_id': item_a}]},
                {'summary': '收现金', 'account_code': '1001', 'debit': 0, 'credit': 10},
            ]})
        self.assertEqual(sc, 400)
        # 回读携带 aux 名称
        _, gv = self._call('get', f'/api/ledger/vouchers/{vid}')
        self.assertEqual(gv['success'], True)
        self.assertEqual(gv['entries'][0]['aux'][0]['item_name'], '公司A')
        # 审核
        sc, d = self._call('post', f'/api/ledger/vouchers/{vid}/audit', json={'action': 'audit'})
        self.assertEqual(sc, 200)
        self.assertEqual(d['success'], True)
        # 新建一张作废
        sc, v2 = self._call('post', '/api/ledger/vouchers', json={
            'vdate': '2099-01-06', 'summary': '待作废',
            'entries': [{'summary': 'a', 'account_code': '1001', 'debit': 20, 'credit': 0},
                        {'summary': 'b', 'account_code': '1001', 'debit': 0, 'credit': 20}]})
        self.assertEqual(sc, 200)
        v2id: int = v2['id']
        sc, _ = self._call('post', f'/api/ledger/vouchers/{v2id}/audit', json={'action': 'void'})
        self.assertEqual(sc, 200)
        _, gv2 = self._call('get', f'/api/ledger/vouchers/{v2id}')
        self.assertEqual(gv2['data']['status'], 'voided')
        # 红冲审核过的凭证
        sc, rev = self._call('post', f'/api/ledger/vouchers/{vid}/reverse')
        self.assertEqual(sc, 200)
        self.assertEqual(rev['success'], True)
        rev_id: int = rev['id']
        # 清理
        for xid in (vid, v2id, rev_id):
            self._call('delete', f'/api/ledger/vouchers/{xid}')
        ar = self._acct('1122')
        self._call('put', f"/api/ledger/accounts/{ar['id']}", json={
            'code': '1122', 'name': ar['name'], 'category': ar['category'],
            'parent_code': ar['parent_code'], 'aux_dims': [],
            'opening': ar.get('opening') or 0, 'remark': ar.get('remark') or ''})
        self._del_dim(dim_id)

    def test_template_crud(self) -> None:
        """凭证模板创建/查询/删除"""
        sc, d = self._call('post', '/api/ledger/templates', json={
            'name': '测试模板',
            'content': [{'summary': '借', 'account_code': '1002', 'debit': 0, 'credit': 0},
                        {'summary': '贷', 'account_code': '5001', 'debit': 0, 'credit': 0}]})
        self.assertEqual(sc, 200)
        self.assertEqual(d['success'], True)
        _, ts = self._call('get', '/api/ledger/templates')
        tmpl = [t for t in ts['data'] if t['name'] == '测试模板'][0]
        self.assertEqual(len(tmpl['content']), 2)
        sc, _ = self._call('delete', f"/api/ledger/templates/{tmpl['id']}")
        self.assertEqual(sc, 200)
        _, ts = self._call('get', '/api/ledger/templates')
        self.assertNotIn('测试模板', [t['name'] for t in ts['data']])

    def test_carry_profit_and_close(self) -> None:
        """结转损益 + 结账锁定 + 反结账"""
        # 动态取收入/费用科目
        _, accts = self._call('get', '/api/ledger/accounts')
        inc_code = [a['code'] for a in accts['data'] if a['category'] == 'income'][0]
        exp_code = [a['code'] for a in accts['data'] if a['category'] == 'expense'][0]
        # 造收入与费用凭证（含借银行存款平衡）
        sc, d = self._call('post', '/api/ledger/vouchers', json={
            'vdate': '2099-01-10', 'summary': '损益业务',
            'entries': [
                {'summary': '费用', 'account_code': exp_code, 'debit': 400, 'credit': 0},
                {'summary': '银行存款', 'account_code': '1002', 'debit': 200, 'credit': 0},
                {'summary': '收入', 'account_code': inc_code, 'debit': 0, 'credit': 600},
            ]})
        self.assertEqual(sc, 200)
        vid: int = d['id']
        # 结转损益
        sc, cp = self._call('post', '/api/ledger/carry-profit', json={'month': '2099-01'})
        self.assertEqual(sc, 200)
        self.assertEqual(cp['success'], True)
        # 结账
        sc, d = self._call('post', '/api/ledger/closing/2099-01')
        self.assertEqual(sc, 200)
        self.assertEqual(d['success'], True)
        # 已结账不允许删除凭证
        sc, d = self._call('delete', f'/api/ledger/vouchers/{vid}')
        self.assertEqual(sc, 400)
        # 重复结账被拒
        sc, d = self._call('post', '/api/ledger/closing/2099-01')
        self.assertEqual(sc, 400)
        self.assertIn('已结账', d.get('error') or '')
        # 反结账
        sc, d = self._call('delete', '/api/ledger/closing/2099-01')
        self.assertEqual(sc, 200)
        self.assertEqual(d['success'], True)
        # 清理
        sc, d = self._call('delete', f'/api/ledger/carry-profit?month=2099-01')
        self.assertEqual(sc, 200)
        self._call('delete', f'/api/ledger/vouchers/{vid}')

    def test_periods_listing(self) -> None:
        """期间列表接口可用"""
        sc, d = self._call('get', '/api/ledger/periods?year=2099')
        self.assertEqual(sc, 200)
        self.assertEqual(d['success'], True)
        self.assertIsInstance(d.get('data'), list)


class TestPurchase(unittest.TestCase):
    """采购台账模块测试（需要登录会话）"""

    client: FlaskClient = None
    _created_ids: list[int] = []

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True
        with cls.client.session_transaction() as sess:  # type: ignore[attr-defined]
            sess['auth'] = True
        cls._rate_json: int = app_mod.RATE_LIMIT_MAX_JSON
        cls._rate_html: int = app_mod.RATE_LIMIT_MAX_HTML
        app_mod.RATE_LIMIT_MAX_JSON = 100000
        app_mod.RATE_LIMIT_MAX_HTML = 100000

    @classmethod
    @override
    def tearDownClass(cls) -> None:
        # 注意顺序：先清数据，再恢复限流。反过来的话清理请求会被限流拦掉（429），
        # 测试数据就残留在真实库里，下一轮跑测试会撞「名称已存在」。
        for rid in cls._created_ids:
            cls.client.delete(f'/api/purchase/records/{rid}')
        app_mod.RATE_LIMIT_MAX_JSON = cls._rate_json
        app_mod.RATE_LIMIT_MAX_HTML = cls._rate_html

    def _create(self, payload: dict) -> dict:
        r = self.client.post('/api/purchase/records', json=payload)
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        data = r.get_json()['data']
        type(self)._created_ids.append(data['id'])
        return data

    def test_page_accessible(self) -> None:
        r = self.client.get('/purchase')
        self.assertEqual(r.status_code, 200)
        html = r.get_data(as_text=True)
        self.assertIn('采购台账', html)
        self.assertIn('基础配置', html)          # 配置面板入口随页面一起下发

    def test_manifest(self) -> None:
        r = self.client.get('/purchase/manifest.json')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json().get('short_name'), '采购台账')

    def test_crud_and_derived(self) -> None:
        """新增 → 派生字段（不含税成本/未付余额/付款状态）→ 编辑 → 删除"""
        rec = self._create({'date': '2026-10-10', 'bill_no': 'T-001',
                            'supplier': '测试供应商甲', 'category': '测试品类',
                            'summary': '测试摘要', 'invoice_type': '专票',
                            'total_amount': 1000, 'input_tax': 90,
                            'invoiced_amount': 500, 'paid_amount': 0,
                            'pay_account': '测试账户'})
        self.assertEqual(rec['net_cost'], 910.0)
        self.assertEqual(rec['unpaid_balance'], 1000.0)
        self.assertEqual(rec['pay_status'], '未付款')

        rid = rec['id']
        r = self.client.post(f'/api/purchase/records/{rid}',
                             json={'date': '2026-10-10', 'supplier': '测试供应商甲',
                                   'invoice_type': '专票', 'total_amount': 1000,
                                   'input_tax': 90, 'paid_amount': 400})
        self.assertEqual(r.status_code, 200)
        upd = r.get_json()['data']
        self.assertEqual(upd['unpaid_balance'], 600.0)
        self.assertEqual(upd['pay_status'], '部分付款')

        r = self.client.post(f'/api/purchase/records/{rid}',
                             json={'date': '2026-10-10', 'supplier': '测试供应商甲',
                                   'total_amount': 1000, 'paid_amount': 1000})
        self.assertEqual(r.get_json()['data']['pay_status'], '已付清')

        r = self.client.delete(f'/api/purchase/records/{rid}')
        self.assertEqual(r.status_code, 200)
        type(self)._created_ids.remove(rid)

    def test_invalid_params(self) -> None:
        cases: list[dict] = [
            {'date': '2026-10-10'},                                   # 缺供应商
            {'date': '2026-10-10', 'supplier': '某', 'total_amount': 0},   # 金额必须 > 0
            {'date': '2026/10/10', 'supplier': '某', 'total_amount': 10},  # 日期格式
            {'date': '2026-10-10', 'supplier': '某', 'total_amount': 10,
             'input_tax': 99},                                        # 进项税 > 总额
            {'date': '2026-10-10', 'supplier': '某', 'total_amount': 10,
             'paid_amount': -1},                                      # 负数
            {'supplier': '某', 'total_amount': 10},                   # 缺日期
        ]
        for payload in cases:
            r = self.client.post('/api/purchase/records', json=payload)
            self.assertEqual(r.status_code, 400, str(payload))
            self.assertFalse(r.get_json()['success'])

    def test_stats_and_filters(self) -> None:
        a = self._create({'date': '2026-10-11', 'supplier': '测试供应商乙',
                          'category': '测试品类', 'invoice_type': '无票(不抵)',
                          'total_amount': 200.5, 'input_tax': 0, 'paid_amount': 200.5})
        b = self._create({'date': '2026-10-12', 'supplier': '测试供应商乙',
                          'category': '测试品类2', 'invoice_type': '专票',
                          'total_amount': 99.5, 'paid_amount': 0})
        r = self.client.get('/api/purchase/stats?supplier=测试供应商乙')
        self.assertEqual(r.status_code, 200)
        s = r.get_json()['data']
        self.assertEqual(s['count'], 2)
        self.assertEqual(s['total'], 300.0)
        self.assertEqual(s['unpaid'], 99.5)
        self.assertEqual(s['net_cost'], 300.0)
        self.assertEqual(len(s['by_supplier']), 1)

        r = self.client.get('/api/purchase/records?supplier=测试供应商乙&status=已付清')
        ids = [x['id'] for x in r.get_json()['data']]
        self.assertIn(a['id'], ids)
        self.assertNotIn(b['id'], ids)

        r = self.client.get('/api/purchase/export?supplier=测试供应商乙')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.headers.get('Content-Type', ''))
        self.assertEqual(r.get_data()[:2], b'PK')     # 导出的是 xlsx（zip 头）

    def test_pagination(self) -> None:
        """page/page_size 分页；不带参数时保持旧行为（全部返回）"""
        dates = ['2026-10-01', '2026-10-02', '2026-10-03', '2026-10-04', '2026-10-05']
        for d in dates:
            self._create({'date': d, 'supplier': '分页测试供应商', 'total_amount': 10})
        base = 'supplier=分页测试供应商'

        r = self.client.get(f'/api/purchase/records?{base}&page=1&page_size=2')
        d = r.get_json()
        self.assertEqual(d['total'], 5)
        self.assertEqual(d['pages'], 3)
        self.assertEqual([x['date'] for x in d['data']], ['2026-10-05', '2026-10-04'])

        d = self.client.get(f'/api/purchase/records?{base}&page=3&page_size=2').get_json()
        self.assertEqual([x['date'] for x in d['data']], ['2026-10-01'])

        # 翻过头 → 停在最后一页
        d = self.client.get(f'/api/purchase/records?{base}&page=99&page_size=2').get_json()
        self.assertEqual(d['page'], 3)
        self.assertEqual(len(d['data']), 1)

        # 与筛选联动
        d = self.client.get(f'/api/purchase/records?{base}'
                            '&start=2026-10-02&end=2026-10-03&page=1&page_size=2').get_json()
        self.assertEqual(d['total'], 2)
        self.assertEqual(d['pages'], 1)

        # 不带分页参数 → 旧行为，全部返回
        d = self.client.get(f'/api/purchase/records?{base}').get_json()
        self.assertEqual(len(d['data']), 5)
        self.assertIsNone(d['page_size'])

    def test_invoice_tax_auto(self) -> None:
        """票种带税率 → 进项税自动算；价税分离与计算抵扣两种算法不能混"""
        # 专票9%：1090 是价税合计 → 税额 = 1090÷1.09×9% = 90
        d = self._create({'date': '2026-10-01', 'supplier': '税率测试供应商',
                          'invoice_type': '专票9%', 'total_amount': 1090})
        self.assertAlmostEqual(d['input_tax'], 90.0, places=2)
        self.assertAlmostEqual(d['net_cost'], 1000.0, places=2)
        # 农产品收购发票：买价 1000 × 扣除率 9% = 90（是买价乘，不价税分离）
        d = self._create({'date': '2026-10-01', 'supplier': '税率测试供应商',
                          'invoice_type': '农产品收购发票', 'total_amount': 1000})
        self.assertAlmostEqual(d['input_tax'], 90.0, places=2)
        # 无票 / 免税 → 不抵扣
        d = self._create({'date': '2026-10-01', 'supplier': '税率测试供应商',
                          'invoice_type': '免税普票', 'total_amount': 500})
        self.assertEqual(d['input_tax'], 0.0)
        # 税额手填优先（票面和系统算的有差异时以票面为准）
        d = self._create({'date': '2026-10-01', 'supplier': '税率测试供应商',
                          'invoice_type': '专票13%', 'total_amount': 1130,
                          'input_tax': 130})
        self.assertEqual(d['input_tax'], 130.0)
        # 老票种「专票」「普票」自动迁到「税额手填」，税额不动
        d = self._create({'date': '2026-10-01', 'supplier': '税率测试供应商',
                          'invoice_type': '专票', 'total_amount': 200, 'input_tax': 18})
        self.assertEqual(d['invoice_type'], '专票(税额手填)')
        self.assertEqual(d['input_tax'], 18.0)

    def test_return_deduct_payable(self) -> None:
        """生鲜常见退货 / 扣款：应付净额 = 采购额 − 退货 − 扣款"""
        d = self._create({'date': '2026-10-02', 'supplier': '退货测试供应商',
                          'invoice_type': '无票(不抵)', 'total_amount': 1000,
                          'return_amount': 100, 'deduct_amount': 30, 'paid_amount': 500})
        self.assertAlmostEqual(d['payable'], 870.0, places=2)
        self.assertAlmostEqual(d['unpaid_balance'], 370.0, places=2)
        self.assertEqual(d['pay_status'], '部分付款')

        # 已付金额不能大于应付净额（原来是跟采购总额比，退货后会对不上）
        r = self.client.post('/api/purchase/records', json={
            'date': '2026-10-03', 'supplier': '退货测试供应商', 'invoice_type': '无票(不抵)',
            'total_amount': 1000, 'deduct_amount': 300, 'paid_amount': 800})
        self.assertEqual(r.status_code, 400)
        self.assertIn('应付净额', r.get_json()['error'])

        # 退货 + 扣款不能大于采购总额
        r = self.client.post('/api/purchase/records', json={
            'date': '2026-10-03', 'supplier': '退货测试供应商', 'invoice_type': '无票(不抵)',
            'total_amount': 1000, 'return_amount': 700, 'deduct_amount': 400})
        self.assertEqual(r.status_code, 400)
        self.assertIn('退货金额', r.get_json()['error'])

    def test_customer_dimension(self) -> None:
        """客户/门店维度：落库 + 按客户统计 + 按客户筛选"""
        self._create({'date': '2026-10-04', 'supplier': '客户测试供应商',
                      'invoice_type': '无票(不抵)', 'total_amount': 100,
                      'customer': '一中食堂'})
        self._create({'date': '2026-10-04', 'supplier': '客户测试供应商',
                      'invoice_type': '无票(不抵)', 'total_amount': 60,
                      'customer': '一中食堂'})
        self._create({'date': '2026-10-04', 'supplier': '客户测试供应商',
                      'invoice_type': '无票(不抵)', 'total_amount': 40,
                      'customer': '二中食堂'})
        d = self.client.get('/api/purchase/stats').get_json()['data']
        rows = {x['customer']: x for x in d['by_customer'] if x['customer']}
        self.assertIn('一中食堂', rows)
        self.assertAlmostEqual(rows['一中食堂']['total'], 160.0, places=2)
        self.assertAlmostEqual(rows['二中食堂']['total'], 40.0, places=2)

        r = self.client.get('/api/purchase/records?customer=一中食堂').get_json()
        self.assertEqual(r['total'], 2)
        self.assertTrue(all(x['customer'] == '一中食堂' for x in r['data']))

        # 客户要能进字典（配置里维护）
        r = self.client.post('/api/purchase/dicts', json={'kind': 'customer', 'name': '三中食堂'})
        self.assertEqual(r.status_code, 200)
        did = r.get_json()['data']['id']
        try:
            opts = self.client.get('/api/purchase/options').get_json()['data']
            self.assertIn('三中食堂', opts['customers'])
        finally:
            self.client.delete(f'/api/purchase/dicts/{did}')

    def test_options(self) -> None:
        r = self.client.get('/api/purchase/options')
        self.assertEqual(r.status_code, 200)
        d = r.get_json()['data']
        self.assertIn('无票(不抵)', d['invoice_types'])

    def test_recon_multi_dimension(self) -> None:
        """对账单多维度：筛选（品类/客户/票种/状态/关键词）+ 汇总维度分组小计"""
        base = {'date': '2026-10-06', 'supplier': '多维对账供应商'}
        self._create(dict(base, category='蔬菜', customer='一中食堂',
                          invoice_type='专票9%', total_amount=1090, paid_amount=0))
        self._create(dict(base, category='肉类', customer='二中食堂',
                          invoice_type='无票(不抵)', total_amount=200, paid_amount=200))
        url = '/api/purchase/reconcile?supplier=多维对账供应商&start=2026-10-01&end=2026-10-31'

        # 汇总维度：按品类 / 按客户 各出两组小计；group_by=none 只出明细
        d = self.client.get(url + '&group_by=category').get_json()['data']
        cats = {g['name']: g for g in d['by_group']}
        self.assertEqual(set(cats), {'蔬菜', '肉类'})
        self.assertAlmostEqual(cats['蔬菜']['total'], 1090.0, places=2)
        self.assertAlmostEqual(cats['肉类']['unpaid'], 0.0, places=2)
        d2 = self.client.get(url + '&group_by=customer').get_json()['data']
        self.assertEqual({g['name'] for g in d2['by_group']}, {'一中食堂', '二中食堂'})
        d3 = self.client.get(url + '&group_by=none').get_json()['data']
        self.assertEqual(d3['by_group'], [])
        self.assertEqual(len(d3['items']), 2)

        # 筛选维度：只对「蔬菜」出单 → 明细、汇总、回显口径都跟着变
        d4 = self.client.get(url + '&category=蔬菜').get_json()['data']
        self.assertEqual(len(d4['items']), 1)
        self.assertAlmostEqual(d4['summary']['total'], 1090.0, places=2)
        self.assertEqual(d4['filters']['category'], '蔬菜')
        # 状态维度：未付款只剩蔬菜那笔（肉类已付清）
        d5 = self.client.get(url + '&status=未付款').get_json()['data']
        self.assertEqual(len(d5['items']), 1)
        self.assertEqual(d5['items'][0]['category'], '蔬菜')
        # 关键词维度
        d6 = self.client.get(url + '&search=二中').get_json()['data']
        self.assertEqual(len(d6['items']), 1)
        self.assertEqual(d6['items'][0]['customer'], '二中食堂')

        # 按日汇总：同一天的两笔合成一行，且按日期升序
        d7 = self.client.get(url + '&group_by=day').get_json()['data']
        days = d7['by_group']
        self.assertEqual(len(days), 1)
        self.assertEqual(days[0]['name'], '2026-10-06')
        self.assertEqual(days[0]['cnt'], 2)
        self.assertAlmostEqual(days[0]['total'], 1290.0, places=2)
        self.assertAlmostEqual(days[0]['payable'], 1290.0, places=2)
        self.assertAlmostEqual(days[0]['unpaid'], 1090.0, places=2)
        # 明细粒度 detail=day 时强制按日期分组
        d8 = self.client.get(url + '&detail=day&group_by=none').get_json()['data']
        self.assertEqual(d8['group_by'], 'day')
        self.assertEqual([g['name'] for g in d8['by_group']], ['2026-10-06'])

        # 导出：按日汇总版 / 只出汇总版 都要能正常出 xlsx
        for q in (url + '&detail=day', url + '&detail=none', url):
            r = self.client.get('/api/purchase/reconcile/export?' + q.replace(
                '/api/purchase/reconcile?', ''))
            self.assertEqual(r.status_code, 200, q)
            self.assertEqual(r.get_data()[:2], b'PK')

    def test_multi_select_filters(self) -> None:
        """分析支持多选：多个供应商 / 多个具体日期可以同时筛，并能叠加"""
        for day, sup, amt in (('2026-10-06', '多选甲', 100), ('2026-10-07', '多选乙', 200),
                              ('2026-10-08', '多选丙', 300)):
            self._create({'date': day, 'supplier': sup, 'total_amount': amt})

        # 多选供应商（逗号分隔）
        d = self.client.get('/api/purchase/stats?supplier=多选甲,多选乙').get_json()['data']
        self.assertEqual({x['supplier'] for x in d['by_supplier']}, {'多选甲', '多选乙'})
        self.assertEqual(d['count'], 2)
        # 多选具体日期
        d2 = self.client.get('/api/purchase/stats?dates=2026-10-06,2026-10-08').get_json()['data']
        self.assertEqual(d2['count'], 2)
        self.assertAlmostEqual(d2['total'], 400.0, places=2)
        # 两个维度叠加：供应商 ∩ 日期
        d3 = self.client.get(
            '/api/purchase/stats?supplier=多选甲,多选丙&dates=2026-10-06').get_json()['data']
        self.assertEqual(d3['count'], 1)
        self.assertEqual(d3['by_supplier'][0]['supplier'], '多选甲')
        # 日期不合法 → 忽略，退回起止区间
        d4 = self.client.get(
            '/api/purchase/stats?dates=xxx&start=2026-10-06&end=2026-10-07').get_json()['data']
        self.assertEqual(d4['count'], 2)
        # 按日趋势、明细、导出都吃同一套参数
        t = self.client.get(
            '/api/purchase/trends?group=day&dates=2026-10-06,2026-10-08').get_json()['data']
        self.assertEqual(sorted(r['period'] for r in t['rows']), ['2026-10-06', '2026-10-08'])
        recs = self.client.get('/api/purchase/records?supplier=多选甲,多选乙').get_json()['data']
        self.assertEqual(len(recs), 2)
        self.assertEqual(self.client.get(
            '/api/purchase/export?supplier=多选甲,多选乙').status_code, 200)

    def test_company_settings(self) -> None:
        """对账单抬头设置：保存后能读回（默认值不能丢）"""
        old = self.client.get('/api/purchase/settings').get_json()['data']
        try:
            r = self.client.post('/api/purchase/settings',
                                 json={'company_name': '抬头测试公司', 'settle_day': '15'})
            self.assertEqual(r.status_code, 200)
            d = r.get_json()['data']
            self.assertEqual(d['company_name'], '抬头测试公司')
            self.assertEqual(d['settle_day'], '15')
            self.assertIn('recon_note', d)          # 没提交的键保留默认值
            # 只认已知键，乱传的键不会进库
            r = self.client.post('/api/purchase/settings', json={'hack': 'x'})
            self.assertTrue(r.get_json()['success'])
            self.assertNotIn('hack', r.get_json()['data'])
        finally:
            self.client.post('/api/purchase/settings',
                             json={'company_name': old.get('company_name', ''),
                                   'settle_day': old.get('settle_day', '10')})

    def test_recon_opening_closing_due(self) -> None:
        """对账单按往来账走：期初未付 + 本期应付 − 本期已付 = 期末未付；月结给最迟付款日"""
        # 上一轮跑完供应商删不掉（有采购记录），结算方式会残留，先置空保证可重复跑
        old = next((x for x in self.client.get('/api/purchase/suppliers').get_json()['data']
                    if x['name'] == '对账测试供应商'), None)
        if old:
            self.client.post(f"/api/purchase/suppliers/{old['id']}",
                             json={'name': '对账测试供应商', 'settle_type': ''})
        # 上期（9月）：采购 1000，已付 300 → 未付 700
        self._create({'date': '2026-09-05', 'supplier': '对账测试供应商',
                      'total_amount': 1000, 'paid_amount': 300})
        # 本期（10月）：采购 2000 − 退货 100 − 扣款 50 = 应付 1850，已付 850 → 未付 1000
        self._create({'date': '2026-10-05', 'supplier': '对账测试供应商', 'total_amount': 2000,
                      'return_amount': 100, 'deduct_amount': 50, 'paid_amount': 850})
        url = '/api/purchase/reconcile?supplier=对账测试供应商&start=2026-10-01&end=2026-10-31'
        d = self.client.get(url).get_json()['data']
        self.assertAlmostEqual(d['opening']['unpaid'], 700.0, places=2)
        self.assertEqual(d['opening']['count'], 1)          # 期初那 1 笔
        self.assertAlmostEqual(d['summary']['payable'], 1850.0, places=2)
        self.assertAlmostEqual(d['summary']['unpaid'], 1000.0, places=2)
        self.assertAlmostEqual(d['closing'], 1700.0, places=2)   # 700 + 1000
        self.assertEqual(len(d['items']), 1)                # 明细只含本期
        self.assertEqual(d['due_date'], '')                 # 没设结算方式 → 不给到期日

        # 供应商设成月结 → 最迟付款日 = 次月 settle_day（默认 10）号
        lst = self.client.get('/api/purchase/suppliers').get_json()['data']
        row = next((x for x in lst if x['name'] == '对账测试供应商'), None)
        if row is None:
            row = self.client.post('/api/purchase/suppliers',
                                   json={'name': '对账测试供应商',
                                         'settle_type': '月结'}).get_json()['data']
        else:
            self.client.post(f"/api/purchase/suppliers/{row['id']}",
                             json={'name': '对账测试供应商', 'settle_type': '月结'})
        try:
            d2 = self.client.get(url).get_json()['data']
            self.assertEqual(d2['settle_type'], '月结')
            self.assertEqual(d2['due_date'], '2026-11-10')
            self.assertAlmostEqual(d2['closing'], 1700.0, places=2)
        finally:
            # 有采购记录的供应商删不掉（后端限制），失败的清理不影响断言
            self.client.delete(f"/api/purchase/suppliers/{row['id']}")

    def test_anonymous_blocked(self) -> None:
        anon = app.test_client()
        self.assertNotEqual(
            anon.post('/api/purchase/records', json={'date': '2026-10-10',
                                                     'supplier': 'x',
                                                     'total_amount': 1}).status_code, 200)


class TestPurchaseConfig(unittest.TestCase):
    """采购模块的基础配置：供应商 / 品类 / 付款账户（需要登录会话）"""

    client: FlaskClient = None
    _supplier_ids: list[int] = []
    _dict_ids: list[int] = []
    _record_ids: list[int] = []

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True
        with cls.client.session_transaction() as sess:  # type: ignore[attr-defined]
            sess['auth'] = True
        cls._rate_json: int = app_mod.RATE_LIMIT_MAX_JSON
        cls._rate_html: int = app_mod.RATE_LIMIT_MAX_HTML
        app_mod.RATE_LIMIT_MAX_JSON = 100000
        app_mod.RATE_LIMIT_MAX_HTML = 100000
        cls._clean_leftovers()

    @classmethod
    def _clean_leftovers(cls) -> None:
        """清掉上一轮可能没删干净的本用例数据（否则撞 UNIQUE 名称，用例直接失败）"""
        from urllib.parse import quote
        for s in cls.client.get('/api/purchase/suppliers').get_json()['data']:
            if not s['name'].startswith('配置测试供应商'):
                continue
            for r in cls.client.get(
                    '/api/purchase/records?search=' + quote(s['name'])).get_json()['data']:
                if r['supplier'] == s['name']:
                    cls.client.delete(f'/api/purchase/records/{r["id"]}')
            cls.client.delete(f'/api/purchase/suppliers/{s["id"]}')
        for kind, items in cls.client.get('/api/purchase/dicts').get_json()['data'].items():
            for d in items:
                if d['name'].startswith('配置测试'):
                    cls.client.delete(f'/api/purchase/dicts/{d["id"]}')

    @classmethod
    @override
    def tearDownClass(cls) -> None:
        # 同上：先清数据，再恢复限流，否则清理请求会被限流拦掉留下脏数据
        for rid in cls._record_ids:
            cls.client.delete(f'/api/purchase/records/{rid}')
        for did in cls._dict_ids:
            cls.client.delete(f'/api/purchase/dicts/{did}')
        for sid in cls._supplier_ids:
            cls.client.delete(f'/api/purchase/suppliers/{sid}')
        app_mod.RATE_LIMIT_MAX_JSON = cls._rate_json
        app_mod.RATE_LIMIT_MAX_HTML = cls._rate_html

    def _supplier(self, payload: dict) -> dict:
        r = self.client.post('/api/purchase/suppliers', json=payload)
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        data = r.get_json()['data']
        type(self)._supplier_ids.append(data['id'])
        return data

    def test_supplier_crud(self) -> None:
        """新增 → 重名拒绝 → 编辑 → 停用 → 无记录可删"""
        a = self._supplier({'name': '配置测试供应商A', 'contact': '张老板',
                            'phone': '13000000000', 'category': '配置测试品类',
                            'default_invoice': '专票', 'note': '测试备注'})
        self.assertEqual(a['category'], '配置测试品类')
        self.assertEqual(a['used_cnt'], 0)

        r = self.client.post('/api/purchase/suppliers', json={'name': '配置测试供应商A'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('已存在', r.get_json()['error'])

        r = self.client.post('/api/purchase/suppliers', json={'name': ''})
        self.assertEqual(r.status_code, 400)

        r = self.client.post(f'/api/purchase/suppliers/{a["id"]}',
                             json={'name': '配置测试供应商A', 'contact': '李老板',
                                   'default_invoice': '无票(不抵)'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['data']['contact'], '李老板')

        # 停用后 active=0，且仍查得到
        r = self.client.post(f'/api/purchase/suppliers/{a["id"]}',
                             json={'name': '配置测试供应商A', 'active': 0})
        self.assertEqual(r.get_json()['data']['active'], 0)
        rows = self.client.get('/api/purchase/suppliers').get_json()['data']
        a_row = next(x for x in rows if x['id'] == a['id'])
        self.assertEqual(a_row['active'], 0)
        self.assertEqual(len(self.client.get(
            '/api/purchase/suppliers?active_only=1').get_json()['data']),
            len(rows) - 1)

        # 没有采购记录 → 可以删除
        r = self.client.delete(f'/api/purchase/suppliers/{a["id"]}')
        self.assertEqual(r.status_code, 200)
        type(self)._supplier_ids.remove(a['id'])

    def test_rename_syncs_history(self) -> None:
        """供应商改名并同步时，历史采购记录跟着改；有记录的供应商不允许删除"""
        s = self._supplier({'name': '配置测试供应商B'})
        rec = self.client.post('/api/purchase/records', json={
            'date': '2026-10-15', 'supplier': '配置测试供应商B',
            'total_amount': 100, 'paid_amount': 0})
        rid = rec.get_json()['data']['id']
        type(self)._record_ids.append(rid)

        r = self.client.post(f'/api/purchase/suppliers/{s["id"]}',
                             json={'name': '配置测试供应商B改名', 'sync_history': 1})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['renamed'], 1)
        names = [x['supplier'] for x in
                 self.client.get('/api/purchase/records?search=配置测试供应商').get_json()['data']]
        self.assertIn('配置测试供应商B改名', names)
        self.assertNotIn('配置测试供应商B', names)

        # 有历史记录 → 禁止删除
        r = self.client.delete(f'/api/purchase/suppliers/{s["id"]}')
        self.assertEqual(r.status_code, 400)
        self.assertIn('停用', r.get_json()['error'])

    def test_dicts(self) -> None:
        """品类 / 付款账户字典：新增、重名拒绝、列表、删除"""
        r = self.client.post('/api/purchase/dicts', json={'kind': 'category',
                                                          'name': '配置测试品类X'})
        self.assertEqual(r.status_code, 200)
        did = r.get_json()['data']['id']
        type(self)._dict_ids.append(did)

        r = self.client.post('/api/purchase/dicts', json={'kind': 'category',
                                                          'name': '配置测试品类X'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('已存在', r.get_json()['error'])

        r = self.client.post('/api/purchase/dicts', json={'kind': '不存在的类型',
                                                          'name': 'x'})
        self.assertEqual(r.status_code, 400)

        names = [x['name'] for x in self.client.get(
            '/api/purchase/dicts?kind=category').get_json()['data']]
        self.assertIn('配置测试品类X', names)

        r = self.client.post(f'/api/purchase/dicts/{did}', json={'name': '配置测试品类Y'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['data']['name'], '配置测试品类Y')

        # 配置的品类要出现在表单建议里
        opts = self.client.get('/api/purchase/options').get_json()['data']
        self.assertIn('配置测试品类Y', opts['categories'])

        r = self.client.delete(f'/api/purchase/dicts/{did}')
        self.assertEqual(r.status_code, 200)
        type(self)._dict_ids.remove(did)

    def test_supplier_profile_fields(self) -> None:
        """供应商档案字段：编码（自动/自定义/冲突）、类型、市场、地址"""
        r = self.client.post('/api/purchase/suppliers', json={
            'name': '配置测试供应商档案A', 's_type': '个体', 'market': '恩施华硒',
            'address': '旗峰大道1号', 'contact': '谭博', 'phone': '13986867680',
            'category': '生鲜'})
        self.assertEqual(r.status_code, 200)
        a = r.get_json()['data']
        type(self)._supplier_ids.append(a['id'])
        self.assertRegex(a['code'], r'^GYS\d{3,}$')      # 编码留空自动生成
        self.assertEqual(a['s_type'], '个体')
        self.assertEqual(a['market'], '恩施华硒')
        self.assertEqual(a['address'], '旗峰大道1号')

        # 编码冲突
        r = self.client.post('/api/purchase/suppliers', json={
            'name': '配置测试供应商档案B', 'code': a['code']})
        self.assertEqual(r.status_code, 400)
        self.assertIn('编码', r.get_json()['error'])

        # 自定义编码 + 编辑
        b = self._supplier({'name': '配置测试供应商档案B', 'code': 'GYS998'})
        self.assertEqual(b['code'], 'GYS998')
        r = self.client.post(f'/api/purchase/suppliers/{a["id"]}', json={
            'name': '配置测试供应商档案A', 'code': a['code'], 's_type': '批发市场',
            'market': '金子坝', 'active': 1})
        self.assertEqual(r.get_json()['data']['s_type'], '批发市场')
        self.assertEqual(r.get_json()['data']['market'], '金子坝')

        # 编码可搜索
        names = [s['name'] for s in self.client.get(
            '/api/purchase/suppliers?q=GYS998').get_json()['data']]
        self.assertIn('配置测试供应商档案B', names)

    def test_supplier_import(self) -> None:
        """供应商 Excel 导入：模板、新增、重名跳过、覆盖更新、坏文件"""
        r = self.client.get('/api/purchase/suppliers/template')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_data()[:2], b'PK')
        from openpyxl import Workbook, load_workbook
        ws = load_workbook(io.BytesIO(r.get_data())).active
        self.assertEqual(ws.cell(row=1, column=1).value, '供应商编码')
        self.assertEqual(ws.cell(row=1, column=2).value, '供应商名称')

        def build(contact: str) -> bytes:
            wb = Workbook()
            w = wb.active
            w.append(['供应商编码', '供应商名称', '供应商类型', '批发市场名称',
                      '经营地址', '法人/负责人', '联系方式', '主营品类',
                      '默认发票类型', '结算方式', '备注'])
            w.append(['', '配置测试供应商导入A', '个体', '恩施华硒', '旗峰大道1号',
                      contact, '13900000000', '生鲜', '专票9%', '月结', ''])
            w.append(['GYS900', '配置测试供应商导入B', '批发市场', '清源路', '',
                      '李老板', '', '调料', '乱写的票种', '乱写的结算', ''])
            w.append(['', '', '', '', '', '', '', '', '', '', ''])   # 空行应被忽略
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        def upload(payload: bytes, mode: str = 'skip', name: str = 's.xlsx'):
            return self.client.post(
                '/api/purchase/suppliers/import',
                data={'file': (io.BytesIO(payload), name), 'mode': mode},
                content_type='multipart/form-data')

        d = upload(build('王老板')).get_json()['data']
        self.assertEqual(d['added'], 2)
        type(self)._supplier_ids.extend(d['added_ids'])
        rows = {s['name']: s for s in
                self.client.get('/api/purchase/suppliers').get_json()['data']}
        a_row = rows['配置测试供应商导入A']
        self.assertEqual(a_row['contact'], '王老板')
        self.assertEqual(a_row['default_invoice'], '专票9%')
        self.assertEqual(a_row['settle_type'], '月结')       # 结算方式能导入
        self.assertEqual(a_row['s_type'], '个体')            # 新档案字段能导入
        self.assertEqual(a_row['market'], '恩施华硒')
        self.assertEqual(a_row['address'], '旗峰大道1号')
        self.assertRegex(a_row['code'], r'^GYS\d{3,}$')      # 编码留空自动生成
        self.assertEqual(rows['配置测试供应商导入B']['code'], 'GYS900')  # 指定编码生效
        self.assertEqual(rows['配置测试供应商导入B']['default_invoice'], '')  # 非法票种留空
        self.assertEqual(rows['配置测试供应商导入B']['settle_type'], '')     # 非法结算方式留空

        d2 = upload(build('王老板'), 'skip').get_json()['data']
        self.assertEqual(d2['added'], 0)
        self.assertEqual(d2['skipped'], 2)               # 重名默认不动
        rows = {s['name']: s for s in
                self.client.get('/api/purchase/suppliers').get_json()['data']}
        self.assertEqual(rows['配置测试供应商导入A']['contact'], '王老板')

        d3 = upload(build('王老板2'), 'update').get_json()['data']
        self.assertEqual(d3['updated'], 2)
        rows = {s['name']: s for s in
                self.client.get('/api/purchase/suppliers').get_json()['data']}
        self.assertEqual(rows['配置测试供应商导入A']['contact'], '王老板2')

        r = upload(b'x', 'skip', 'a.xls')
        self.assertEqual(r.status_code, 400)
        self.assertIn('.xlsx', r.get_json()['error'])
        r = upload(build('王老板'), 'skip', 'a.xlsx')     # 内容正常但换后缀不影响
        self.assertEqual(r.status_code, 200)

        bad = Workbook()
        bad.active.append(['姓名', '年龄'])
        buf = io.BytesIO()
        bad.save(buf)
        r = upload(buf.getvalue())
        self.assertEqual(r.status_code, 400)
        self.assertIn('表头', r.get_json()['error'])

    def test_config_anonymous_blocked(self) -> None:
        anon = app.test_client()
        self.assertNotEqual(
            anon.post('/api/purchase/suppliers', json={'name': '匿名供应商'}).status_code, 200)
        self.assertNotEqual(
            anon.get('/api/purchase/suppliers').status_code, 200)


class TestPurchaseStatsAndReconcile(unittest.TestCase):
    """统计维度（年/月/日/供应商）与对账报表（需要登录会话）"""

    client: FlaskClient = None
    _record_ids: list[int] = []

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True
        with cls.client.session_transaction() as sess:  # type: ignore[attr-defined]
            sess['auth'] = True
        cls._rate_json: int = app_mod.RATE_LIMIT_MAX_JSON
        cls._rate_html: int = app_mod.RATE_LIMIT_MAX_HTML
        app_mod.RATE_LIMIT_MAX_JSON = 100000
        app_mod.RATE_LIMIT_MAX_HTML = 100000
        cls._clean_leftovers()
        # 三笔：两笔在 10 月（其中一笔部分付款），一笔在 9 月（用于验证「累计未付」）
        for payload in ({'date': '2026-10-05', 'supplier': '对账测试供应商',
                         'category': '对账测试品类', 'invoice_type': '专票',
                         'total_amount': 100, 'input_tax': 9, 'paid_amount': 60},
                        {'date': '2026-10-20', 'supplier': '对账测试供应商',
                         'invoice_type': '无票(不抵)', 'total_amount': 50},
                        {'date': '2026-09-30', 'supplier': '对账测试供应商',
                         'invoice_type': '普票', 'total_amount': 30}):
            r = cls.client.post('/api/purchase/records', json=payload)
            cls._record_ids.append(r.get_json()['data']['id'])

    @classmethod
    def _clean_leftovers(cls) -> None:
        from urllib.parse import quote
        for r in cls.client.get(
                '/api/purchase/records?search=' + quote('对账测试供应商')).get_json()['data']:
            cls.client.delete(f'/api/purchase/records/{r["id"]}')

    @classmethod
    @override
    def tearDownClass(cls) -> None:
        # 先清数据再恢复限流（规范 7.43）
        for rid in cls._record_ids:
            cls.client.delete(f'/api/purchase/records/{rid}')
        app_mod.RATE_LIMIT_MAX_JSON = cls._rate_json
        app_mod.RATE_LIMIT_MAX_HTML = cls._rate_html

    def test_trends_by_month_year_day(self) -> None:
        base = 'supplier=' + '对账测试供应商'
        r = self.client.get('/api/purchase/trends?group=month&' + base)
        self.assertEqual(r.status_code, 200)
        d = r.get_json()['data']
        self.assertEqual(d['group'], 'month')
        rows = {x['period']: x for x in d['rows']}
        self.assertEqual(rows['2026-10']['total'], 150.0)
        self.assertEqual(rows['2026-10']['cnt'], 2)
        self.assertEqual(rows['2026-10']['paid'], 60.0)
        self.assertEqual(rows['2026-10']['unpaid'], 90.0)
        self.assertEqual(rows['2026-09']['total'], 30.0)

        r = self.client.get('/api/purchase/trends?group=year&' + base)
        rows = {x['period']: x for x in r.get_json()['data']['rows']}
        self.assertEqual(rows['2026']['total'], 180.0)
        self.assertEqual(rows['2026']['cnt'], 3)

        r = self.client.get('/api/purchase/trends?group=day&start=2026-10-01&end=2026-10-31&' + base)
        rows = {x['period']: x for x in r.get_json()['data']['rows']}
        self.assertEqual(sorted(rows), ['2026-10-05', '2026-10-20'])
        self.assertEqual(rows['2026-10-05']['total'], 100.0)

        r = self.client.get('/api/purchase/trends?group=week')
        self.assertEqual(r.status_code, 400)

    def test_stats_supplier_has_full_fields(self) -> None:
        s = self.client.get('/api/purchase/stats?supplier=对账测试供应商').get_json()['data']
        row = s['by_supplier'][0]
        self.assertEqual(row['total'], 180.0)
        self.assertEqual(row['tax'], 9.0)
        self.assertEqual(row['net_cost'], 171.0)
        self.assertEqual(row['paid'], 60.0)
        self.assertEqual(row['unpaid'], 120.0)

    def test_reconcile_summary_and_outstanding(self) -> None:
        r = self.client.get('/api/purchase/reconcile?supplier=对账测试供应商'
                            '&start=2026-10-01&end=2026-10-31')
        self.assertEqual(r.status_code, 200)
        d = r.get_json()['data']
        self.assertEqual(d['supplier'], '对账测试供应商')
        self.assertEqual(len(d['items']), 2)
        s = d['summary']
        self.assertEqual(s['count'], 2)
        self.assertEqual(s['total'], 150.0)
        self.assertEqual(s['tax'], 9.0)
        self.assertEqual(s['net_cost'], 141.0)
        self.assertEqual(s['invoiced'], 0.0)
        self.assertEqual(s['uninvoiced'], 150.0)     # 待开票
        self.assertEqual(s['paid'], 60.0)
        self.assertEqual(s['unpaid'], 90.0)
        self.assertEqual(s['recon_cnt'], 0)
        self.assertEqual(s['unrecon_cnt'], 2)
        # 累计未付看全部历史，不只本期间
        self.assertEqual(d['outstanding']['count'], 3)
        self.assertEqual(d['outstanding']['unpaid'], 120.0)
        self.assertEqual(d['by_month'][0]['period'], '2026-10')
        self.assertEqual(d['by_month'][0]['total'], 150.0)
        self.assertEqual(len(d['by_invoice']), 2)

    def test_reconcile_mark_and_export(self) -> None:
        d = self.client.get('/api/purchase/reconcile?supplier=对账测试供应商'
                            '&start=2026-10-01&end=2026-10-31').get_json()['data']
        rid = d['items'][0]['id']
        r = self.client.post(f'/api/purchase/records/{rid}/reconcile', json={'reconciled': 1})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()['data']['reconciled'])

        d2 = self.client.get('/api/purchase/reconcile?supplier=对账测试供应商'
                             '&start=2026-10-01&end=2026-10-31').get_json()['data']
        self.assertEqual(d2['summary']['recon_cnt'], 1)
        self.assertEqual(d2['summary']['unrecon_cnt'], 1)

        # 撤销
        r = self.client.post(f'/api/purchase/records/{rid}/reconcile', json={'reconciled': 0})
        self.assertFalse(r.get_json()['data']['reconciled'])

        xlsx = self.client.get('/api/purchase/reconcile/export?supplier=对账测试供应商'
                               '&start=2026-10-01&end=2026-10-31').get_data()
        self.assertEqual(xlsx[:2], b'PK')              # 对账单同样是 xlsx
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(xlsx)).active
        self.assertEqual(ws.cell(row=1, column=1).value, '采购对账单')
        self.assertEqual(ws.cell(row=2, column=2).value, '对账测试供应商')

        r = self.client.post('/api/purchase/records/999999/reconcile', json={'reconciled': 1})
        self.assertEqual(r.status_code, 404)

    def test_stats_anonymous_blocked(self) -> None:
        anon = app.test_client()
        self.assertNotEqual(anon.get('/api/purchase/trends?group=month').status_code, 200)
        self.assertNotEqual(anon.get('/api/purchase/reconcile').status_code, 200)


class TestPurchaseExcel(unittest.TestCase):
    """采购台账的 Excel 导出 / 模板 / 导入（需要登录会话）"""

    client: FlaskClient = None
    _record_ids: list[int] = []
    _supplier_ids: list[int] = []

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True
        with cls.client.session_transaction() as sess:  # type: ignore[attr-defined]
            sess['auth'] = True
        cls._rate_json: int = app_mod.RATE_LIMIT_MAX_JSON
        cls._rate_html: int = app_mod.RATE_LIMIT_MAX_HTML
        app_mod.RATE_LIMIT_MAX_JSON = 100000
        app_mod.RATE_LIMIT_MAX_HTML = 100000
        cls._clean_leftovers()

    @classmethod
    def _clean_leftovers(cls) -> None:
        from urllib.parse import quote
        for r in cls.client.get(
                '/api/purchase/records?search=' + quote('导入测试')).get_json()['data']:
            cls.client.delete(f'/api/purchase/records/{r["id"]}')
        for s in cls.client.get('/api/purchase/suppliers').get_json()['data']:
            if s['name'].startswith('导入测试'):
                cls.client.delete(f'/api/purchase/suppliers/{s["id"]}')

    @classmethod
    @override
    def tearDownClass(cls) -> None:
        # 先清数据再恢复限流（规范 7.43）
        for rid in cls._record_ids:
            cls.client.delete(f'/api/purchase/records/{rid}')
        for sid in cls._supplier_ids:
            cls.client.delete(f'/api/purchase/suppliers/{sid}')
        app_mod.RATE_LIMIT_MAX_JSON = cls._rate_json
        app_mod.RATE_LIMIT_MAX_HTML = cls._rate_html

    def _xlsx(self, rows: list[list]) -> bytes:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _upload(self, payload: bytes, name: str = 'test.xlsx'):
        return self.client.post(
            '/api/purchase/import',
            data={'file': (io.BytesIO(payload), name), 'auto_supplier': '1'},
            content_type='multipart/form-data')

    def test_template_is_xlsx(self) -> None:
        r = self.client.get('/api/purchase/template')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.headers.get('Content-Type', ''))
        self.assertEqual(r.get_data()[:2], b'PK')          # xlsx 就是 zip
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.get_data())).active
        self.assertEqual(ws.cell(row=1, column=3).value, '供应商名称')

    def test_export_is_xlsx(self) -> None:
        rec = self.client.post('/api/purchase/records', json={
            'date': '2026-10-06', 'supplier': '导入测试导出用',
            'invoice_type': '专票', 'total_amount': 226, 'input_tax': 26,
            'paid_amount': 100})
        type(self)._record_ids.append(rec.get_json()['data']['id'])
        r = self.client.get('/api/purchase/export')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.headers.get('Content-Type', ''))
        body = r.get_data()
        self.assertEqual(body[:2], b'PK')
        # CSV 已弃用：导出不再是纯文本
        self.assertNotIn(b'\xef\xbb\xbf', body[:4])
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(body)).active
        self.assertEqual(ws.cell(row=1, column=1).value, '日期')
        self.assertEqual(ws.freeze_panes, 'A2')            # 表头冻结
        # 第 8 列是采购总金额（1日期 2单据号 3供应商 4客户 5品类 6摘要 7票种）
        self.assertEqual(ws.cell(row=2, column=8).number_format, '#,##0.00')

    def test_import_roundtrip(self) -> None:
        """两行正常 + 一行金额非法 + 一行缺供应商；供应商自动建档"""
        xlsx = self._xlsx([
            ['日期', '供应商名称', '采购总金额', '单据号', '品类', '发票类型', '进项税额'],
            ['2026/10/01', '导入测试供应商A', '¥1,130.00', 'IMP-001', '生鲜', '专票', 130],
            ['2026-10-02', '导入测试供应商B', 500, 'IMP-002', '调料', '无票(不抵)', 0],
            ['2026-10-03', '导入测试供应商C', 'abc', 'IMP-003', '生鲜', '专票', 0],
            ['2026-10-04', '', 100, 'IMP-004', '生鲜', '专票', 0],
        ])
        r = self._upload(xlsx)
        self.assertEqual(r.status_code, 200)
        d = r.get_json()['data']
        self.assertEqual(d['imported'], 2)
        self.assertEqual(d['skipped'], 0)
        self.assertEqual(len(d['errors']), 2)
        self.assertTrue(any('金额' in e['reason'] for e in d['errors']))
        type(self)._record_ids.extend(d['imported_ids'])

        recs = self.client.get('/api/purchase/records?supplier=导入测试供应商A').get_json()['data']
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0]['total_amount'], 1130.0)     # ¥1,130.00 被还原成数字
        self.assertEqual(recs[0]['net_cost'], 1000.0)
        self.assertEqual(recs[0]['date'], '2026-10-01')       # 2026/10/01 被规范成标准格式

        # 新供应商自动进了配置
        names = [s['name'] for s in self.client.get('/api/purchase/suppliers').get_json()['data']]
        self.assertIn('导入测试供应商A', names)
        for s in self.client.get('/api/purchase/suppliers').get_json()['data']:
            if s['name'].startswith('导入测试'):
                type(self)._supplier_ids.append(s['id'])

        # 再导一次：两条都因单据号重复被跳过，不会被覆盖
        d2 = self._upload(xlsx).get_json()['data']
        self.assertEqual(d2['imported'], 0)
        self.assertEqual(d2['skipped'], 2)

    def test_import_judges_by_content_not_suffix(self) -> None:
        """内容是真的 xlsx 就该收：大写后缀 / .xlsm / 文件名末尾有空格 / .csv 后缀"""
        for i, name in enumerate(('Book.XLSX', 'book.xlsm', 'book.xlsx ', 'book.csv')):
            xlsx = self._xlsx([
                ['日期', '供应商名称', '采购总金额', '单据号'],
                ['2026-10-08', '导入测试供应商D', 10, f'IMP-SFX-{i}'],
            ])
            r = self._upload(xlsx, name)
            self.assertEqual(r.status_code, 200, f'{name} 被拒：{r.get_json().get("error")}')
            d = r.get_json()['data']
            self.assertEqual(d['imported'], 1, f'{name} 没导进去：{d}')
            type(self)._record_ids.extend(d['imported_ids'])

    def test_import_rejects_real_xls_and_corrupt(self) -> None:
        """老版 .xls（OLE2 头）、空文件、假 xlsx（普通 zip）都要给明确提示"""
        import zipfile
        r = self._upload(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64, 'legacy.xls')
        self.assertEqual(r.status_code, 400)
        self.assertIn('另存为', r.get_json()['error'])

        r = self._upload(b'', 'empty.xlsx')
        self.assertEqual(r.status_code, 400)

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z:
            z.writestr('hello.txt', 'hi')
        r = self._upload(buf.getvalue(), 'fake.xlsx')     # zip 头对，但不是工作簿
        self.assertEqual(r.status_code, 400)

    def test_import_template_has_no_sample_rows(self) -> None:
        """模板自带的示例行不能被当成真实数据导入"""
        r = self._upload(self.client.get('/api/purchase/template').get_data(),
                         'template.xlsx')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['data']['imported'], 0)

    def test_import_rejects_bad_files(self) -> None:
        r = self._upload(b'not an excel at all', 'a.xls')
        self.assertEqual(r.status_code, 400)
        self.assertIn('.xlsx', r.get_json()['error'])

        r = self._upload(b'hello', 'a.xlsx')               # 后缀对但内容不是 xlsx
        self.assertEqual(r.status_code, 400)

        r = self._upload(self._xlsx([['姓名', '年龄'], ['张三', 18]]))   # 认不出表头
        self.assertEqual(r.status_code, 400)
        self.assertIn('表头', r.get_json()['error'])

        r = self.client.post('/api/purchase/import',
                             data={}, content_type='multipart/form-data')
        self.assertEqual(r.status_code, 400)

    def test_excel_anonymous_blocked(self) -> None:
        anon = app.test_client()
        self.assertNotEqual(anon.get('/api/purchase/export').status_code, 200)
        self.assertNotEqual(anon.get('/api/purchase/template').status_code, 200)
        self.assertNotEqual(
            anon.post('/api/purchase/import',
                      data={'file': (io.BytesIO(b'x'), 'a.xlsx')},
                      content_type='multipart/form-data').status_code, 200)


class TestRateLimiting(unittest.TestCase):
    """速率限制测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_normal_request_passes(self) -> None:
        """正常请求不被限制"""
        r = self.client.get('/api/version')
        self.assertEqual(r.status_code, 200)

    def test_html_page_not_limited_easily(self) -> None:
        """HTML 页面有较高限制"""
        r = self.client.get('/')
        self.assertEqual(r.status_code, 200)


class TestDeployAPI(unittest.TestCase):
    """部署 API 测试"""

    client: FlaskClient = None

    @classmethod
    @override
    def setUpClass(cls) -> None:
        cls.client = app.test_client()
        app.config['TESTING'] = True

    def test_ping_no_token(self) -> None:
        """没有 token 也能 ping"""
        r = self.client.get('/api/ping')
        self.assertEqual(r.status_code, 200)

    def test_pull_without_token(self) -> None:
        """没有 token 应该被拒绝"""
        r = self.client.post('/api/git-pull')
        self.assertEqual(r.status_code, 403)

    def test_reload_without_token(self) -> None:
        """没有 token 应该被拒绝"""
        r = self.client.post('/api/reload-webapp')
        self.assertEqual(r.status_code, 403)


class TestDBConnection(unittest.TestCase):
    """数据库连接测试"""

    def test_sqlite_create_and_write(self) -> None:
        """测试 SQLite WAL 模式连接"""
        with tempfile.NamedTemporaryFile(suffix='.db', delete=False) as f:
            db_path: str = f.name
        try:
            conn = sqlite3.connect(db_path)
            _ = conn.execute('PRAGMA journal_mode=WAL')
            _ = conn.execute('CREATE TABLE test (id INTEGER PRIMARY KEY, name TEXT)')
            _ = conn.execute('INSERT INTO test (name) VALUES (?)', ('hello',))
            conn.commit()
            row = conn.execute('SELECT name FROM test WHERE id=1').fetchone()
            self.assertEqual(row[0], 'hello')
            conn.close()
        finally:
            os.unlink(db_path)


if __name__ == '__main__':
    _ = unittest.main(verbosity=2)
