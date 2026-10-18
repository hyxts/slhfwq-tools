# -*- coding: utf-8 -*-
"""临时：验证 Excel 导出 / 模板 / 导入（含重复跳过、错误行、模板示例行）"""
import io
import sys

sys.path.insert(0, '.')
import app as app_mod

from openpyxl import Workbook, load_workbook

app_mod.app.config['TESTING'] = True
app_mod.RATE_LIMIT_MAX_JSON = 100000
app_mod.RATE_LIMIT_MAX_HTML = 100000
c = app_mod.app.test_client()
with c.session_transaction() as s:
    s['auth'] = True

# 1) 模板
r = c.get('/api/purchase/template')
print('template', r.status_code, r.headers.get('Content-Type'))
wb = load_workbook(io.BytesIO(r.get_data()))
print('template sheet cols:', [cell.value for cell in wb.active[1]][:4], '...')

# 2) 造一个待导入的 xlsx：2 行正常 + 1 行金额非法 + 1 行缺供应商
ws_wb = Workbook()
ws = ws_wb.active
ws.append(['日期', '供应商名称', '采购总金额', '单据号', '品类', '发票类型', '进项税额'])
ws.append(['2026/10/01', '导入测试供应商A', '¥1,130.00', 'IMP-001', '生鲜', '专票', 130])
ws.append(['2026-10-02', '导入测试供应商B', 500, 'IMP-002', '调料', '无票(不抵)', 0])
ws.append(['2026-10-03', '导入测试供应商C', 'abc', 'IMP-003', '生鲜', '专票', 0])
ws.append(['2026-10-04', '', 100, 'IMP-004', '生鲜', '专票', 0])
buf = io.BytesIO()
ws_wb.save(buf)
data = buf.getvalue()

r = c.post('/api/purchase/import',
           data={'file': (io.BytesIO(data), 'test.xlsx'), 'auto_supplier': '1'},
           content_type='multipart/form-data')
print('import', r.status_code, r.get_json())

# 3) 再导一次（应全部跳过重复）
r2 = c.post('/api/purchase/import',
            data={'file': (io.BytesIO(data), 'test.xlsx'), 'auto_supplier': '1'},
            content_type='multipart/form-data')
print('import again', r2.get_json()['data']['imported'], 'skipped',
      r2.get_json()['data']['skipped'])

# 4) 导出台账 xlsx，检查文件内容
r = c.get('/api/purchase/export')
print('export', r.status_code, r.headers.get('Content-Type'))
wbx = load_workbook(io.BytesIO(r.get_data()))
print('export rows:', wbx.active.max_row, 'freeze:', wbx.active.freeze_panes)

# 5) 导入模板自带的示例行不应被导入
tpl = io.BytesIO()
wb.save(tpl)
r = c.post('/api/purchase/import',
           data={'file': (io.BytesIO(tpl.getvalue()), 'template.xlsx')},
           content_type='multipart/form-data')
d = r.get_json()
print('import template ->', d.get('success'), 'imported:', d.get('data', {}).get('imported') if d.get('success') else d.get('error'))

# 6) 非 xlsx 文件应被拒
r = c.post('/api/purchase/import',
           data={'file': (io.BytesIO(b'hello'), 'a.xls')},
           content_type='multipart/form-data')
print('bad file', r.status_code, r.get_json().get('error'))

# 清理
for s in ('导入测试供应商A', '导入测试供应商B'):
    for rec in c.get('/api/purchase/records?supplier=' + s).get_json()['data']:
        c.delete(f"/api/purchase/records/{rec['id']}")
    rows = c.get('/api/purchase/suppliers').get_json()['data']
    for row in rows:
        if row['name'] == s:
            c.delete(f"/api/purchase/suppliers/{row['id']}")
print('cleanup done')
